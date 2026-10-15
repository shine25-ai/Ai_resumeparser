"""
Interview service handling interview scheduling, rescheduling, feedback, filtering, and status updates.
"""

import csv
import html
import io
import smtplib
import uuid
import re
import urllib.parse
from email.message import EmailMessage
from pathlib import Path
from typing import Any, Dict, List, Optional
from fastapi import HTTPException, UploadFile
from loguru import logger
from app.core.exceptions import NotFoundError
from app.models.interview import InterviewDocument
from app.repositories.interview_repository import InterviewRepository
from app.schemas.interview import (
    BulkInterviewFeedbackRequest,
    CandidateFullHistoryResponse,
    InterviewBatchCreateRequest,
    InterviewCancelRequest,
    InterviewCheckConflictRequest,
    InterviewCheckConflictResponse,
    InterviewCreateRequest,
    InterviewFeedbackRequest,
    InterviewListResponse,
    InterviewRescheduleRequest,
    InterviewResponse,
    InterviewUpdateRequest,
    SendInterviewEmailRequest,
    ShareInterviewSummaryRequest,
)
from app.services.s3_service import S3Service
from app.services.settings_service import SettingsService
from app.services.template_service import TemplateService
from app.utils.encryption import decrypt_password
from app.utils.enums import InterviewStatus
from app.utils.helpers import utc_now
from app.utils.validators import validate_uploaded_file


def build_summary_csv(headers: List[str], rows: List[List[Any]]) -> bytes:
    """Build the interview summary CSV attachment."""

    def csv_safe(value: Any) -> str:
        # Bullet lines ("- point") would otherwise open as broken formulas in Excel
        text = re.sub(r"(?m)^- ", "• ", str(value)) if value not in (None, "") else ""
        return f"'{text}" if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows([[csv_safe(v) for v in row] for row in rows])
    # utf-8-sig adds a BOM so Excel opens non-ASCII names correctly
    return buffer.getvalue().encode("utf-8-sig")


def build_summary_xlsx(headers: List[str], rows: List[List[Any]], status_colors: Dict[str, tuple]) -> bytes:
    """Build the interview summary Excel attachment: bold frozen header, filters, wrapped feedback, coloured status."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as err:
        raise HTTPException(status_code=500, detail="Excel export is unavailable: install the 'openpyxl' package on the server.") from err

    wb = Workbook()
    ws = wb.active
    ws.title = "Interview Status"
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="4F46E5")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)

    status_col = headers.index("Status") + 1
    rating_col = headers.index("Rating") + 1
    for row in rows:
        values: List[Any] = []
        for v in row:
            if v in (None, ""):
                values.append(None)
            elif isinstance(v, (int, float)):
                values.append(v)
            else:
                values.append(str(v))
        ws.append(values)
        r = ws.max_row
        for c in ws[r]:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            # openpyxl stores strings starting with "=" as formulas; keep them as plain text
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"
        rating_cell = ws.cell(row=r, column=rating_col)
        try:
            rating_cell.value = float(rating_cell.value) if rating_cell.value is not None else None
        except (TypeError, ValueError):
            pass
        status_cell = ws.cell(row=r, column=status_col)
        colors = status_colors.get(str(status_cell.value or ""))
        if colors:
            bg, fg = colors
            status_cell.fill = PatternFill("solid", fgColor=bg.lstrip("#").upper())
            status_cell.font = Font(bold=True, color=fg.lstrip("#").upper())

    widths = {
        "Candidate Name": 24, "Candidate Email": 28, "Job Title": 22, "Round": 20, "Round Number": 9,
        "Scheduled Date": 14, "Scheduled Time": 11, "Interviewer": 20, "Status": 20, "Rating": 8,
        "Outcome": 16, "Feedback Comments": 50, "Client Feedback Comments": 50,
    }
    for idx, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(h, 16)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class InterviewService:
    """Service handling interview business logic."""

    def __init__(self, interview_repo: InterviewRepository, s3_service: Optional[S3Service] = None):
        self.interview_repo = interview_repo
        self.s3_service = s3_service or S3Service()

    async def check_interview_conflict(self, payload: InterviewCheckConflictRequest) -> InterviewCheckConflictResponse:
        """Check if any interviewer or client has a time slot conflict for the requested date and time."""
        res = await self.interview_repo.check_schedule_conflict(
            scheduled_date=payload.scheduled_date,
            scheduled_time=payload.scheduled_time,
            interviewer_id=payload.interviewer_id,
            interviewer_name=payload.interviewer_name,
            client_id=payload.client_id,
            client_name=payload.client_name,
            interviewers=[i.model_dump() for i in payload.interviewers] if payload.interviewers else None,
            clients=[c.model_dump() for c in payload.clients] if payload.clients else None,
            exclude_interview_id=payload.exclude_interview_id,
        )
        return InterviewCheckConflictResponse(
            has_conflict=res["has_conflict"],
            conflict_type=res.get("conflict_type"),
            conflict_message=res.get("conflict_message"),
            conflicting_interviews=res.get("conflicting_interviews", []),
        )

    async def upload_interview_document(self, file: UploadFile) -> dict:
        """Upload an interview document file to AWS S3 and return the public URL."""
        content = await validate_uploaded_file(file)
        ext = Path(file.filename or "file").suffix
        safe_filename = file.filename or f"document{ext}"
        unique_filename = f"{uuid.uuid4().hex[:12]}_{safe_filename}"
        s3_url = self.s3_service.upload_file(
            file_content=content,
            filename=unique_filename,
            content_type=file.content_type or "application/octet-stream",
            folder="interview_documents",
        )
        logger.info(f"Uploaded interview document to S3: {s3_url}")
        return {
            "s3_url": s3_url,
            "filename": unique_filename,
            "original_filename": safe_filename,
        }

    async def create_interview(self, payload: InterviewCreateRequest, created_by: Optional[str] = None) -> InterviewResponse:
        """Schedule a new interview document."""
        # Validation: Ensure candidate does not already have an active incomplete interview session
        active_interviews = await self.interview_repo.get_active_incomplete_by_candidate(
            candidate_id=payload.candidate_id, candidate_name=payload.candidate_name
        )
        if active_interviews:
            active = active_interviews[0]
            st = active.get("status", "SCHEDULED")
            raise HTTPException(
                status_code=400,
                detail=f"Cannot schedule interview/next round. Candidate '{payload.candidate_name}' already has an active interview session (Status: {st}) that is not completed yet. Please submit feedback to complete the previous round first."
            )

        # Validation: Check for interviewer/client schedule time conflict
        conflict_res = await self.interview_repo.check_schedule_conflict(
            scheduled_date=payload.scheduled_date,
            scheduled_time=payload.scheduled_time,
            interviewer_id=payload.interviewer_id,
            interviewer_name=payload.interviewer_name,
            client_id=payload.client_id,
            client_name=payload.client_name,
            interviewers=[i.model_dump() for i in payload.interviewers] if payload.interviewers else None,
            clients=[c.model_dump() for c in payload.clients] if payload.clients else None,
        )
        if conflict_res.get("has_conflict"):
            raise HTTPException(status_code=400, detail=conflict_res.get("conflict_message"))

        interview_doc = InterviewDocument(
            candidate_id=payload.candidate_id,
            candidate_name=payload.candidate_name,
            candidate_email=payload.candidate_email,
            resume_id=payload.resume_id,
            job_id=payload.job_id,
            job_title=payload.job_title,
            job_location=payload.job_location,
            job_type=payload.job_type,
            interview_type=payload.interview_type,
            interview_type_id=payload.interview_type_id,
            round_number=payload.round_number,
            scheduled_date=payload.scheduled_date,
            scheduled_time=payload.scheduled_time,
            timezone=payload.timezone,
            duration_minutes=payload.duration_minutes,
            interviewer_id=payload.interviewer_id,
            interviewer_name=payload.interviewer_name,
            interviewer_email=payload.interviewer_email,
            meeting_link=payload.meeting_link,
            meeting_platform=payload.meeting_platform,
            location=payload.location or payload.interview_location,
            interview_location=payload.interview_location or payload.location,
            hr_call_verification=payload.hr_call_verification,
            candidate_requested_date_time=payload.candidate_requested_date_time,
            candidate_requested_date=payload.candidate_requested_date,
            candidate_requested_time=payload.candidate_requested_time,
            candidate_requested_role=payload.candidate_requested_role,
            salary_requested=payload.salary_requested,
            final_fit_salary=payload.final_fit_salary,
            joining_date=payload.joining_date,
            interview_document_files=payload.interview_document_files,
            recommendation=payload.recommendation or "Pending",
            status=InterviewStatus.SCHEDULED,
            client_rating=payload.client_rating,
            client_feedback=payload.client_feedback,
            client_strengths=payload.client_strengths,
            client_weaknesses=payload.client_weaknesses,
            client_recommendation=payload.client_recommendation,
            client_notes=payload.client_notes,
            client_id=payload.client_id,
            client_name=payload.client_name,
            client_email=payload.client_email,
            client_feedback_date=payload.client_feedback_date,
            interviewers=[i.model_dump() for i in payload.interviewers] if payload.interviewers else [],
            clients=[c.model_dump() for c in payload.clients] if payload.clients else [],
            notes=payload.notes,
            created_by=created_by,
            updated_by=created_by,
        )

        created = await self.interview_repo.create(interview_doc.to_dict())
        logger.info(f"Scheduled new interview ID '{interview_doc.id}' for candidate '{payload.candidate_name}'")
        return InterviewResponse.model_validate(created)

    async def batch_create_interviews(
        self, payload: InterviewBatchCreateRequest, created_by: Optional[str] = None
    ) -> List[InterviewResponse]:
        """Batch schedule interviews for multiple candidates globally."""
        # Validation: Check for interviewer/client schedule time conflict
        conflict_res = await self.interview_repo.check_schedule_conflict(
            scheduled_date=payload.scheduled_date,
            scheduled_time=payload.scheduled_time,
            interviewer_id=payload.interviewer_id,
            interviewer_name=payload.interviewer_name,
            client_id=payload.client_id,
            client_name=payload.client_name,
            interviewers=[i.model_dump() for i in payload.interviewers] if payload.interviewers else None,
            clients=[c.model_dump() for c in payload.clients] if payload.clients else None,
        )
        if conflict_res.get("has_conflict"):
            raise HTTPException(status_code=400, detail=conflict_res.get("conflict_message"))

        created_interviews: List[InterviewResponse] = []

        for candidate in payload.candidates:
            # Validation: Ensure candidate does not already have an active incomplete interview session
            active_interviews = await self.interview_repo.get_active_incomplete_by_candidate(
                candidate_id=candidate.candidate_id, candidate_name=candidate.candidate_name
            )
            if active_interviews:
                active = active_interviews[0]
                st = active.get("status", "SCHEDULED")
                raise HTTPException(
                    status_code=400,
                    detail=f"Cannot batch schedule interview. Candidate '{candidate.candidate_name}' already has an active interview session (Status: {st}) that is not completed yet. Please complete current interview round first."
                )

            loc = candidate.interview_location or candidate.location or payload.interview_location or payload.location
            interview_doc = InterviewDocument(
                candidate_id=candidate.candidate_id,
                candidate_name=candidate.candidate_name,
                resume_id=candidate.resume_id,
                job_id=payload.job_id,
                job_title=payload.job_title,
                job_location=payload.job_location,
                job_type=payload.job_type,
                interview_type=payload.interview_type,
                interview_type_id=payload.interview_type_id,
                round_number=payload.round_number,
                scheduled_date=payload.scheduled_date,
                scheduled_time=payload.scheduled_time,
                timezone=payload.timezone,
                duration_minutes=payload.duration_minutes,
                interviewer_id=payload.interviewer_id,
                interviewer_name=payload.interviewer_name,
                interviewer_email=payload.interviewer_email,
                meeting_link=payload.meeting_link,
                meeting_platform=payload.meeting_platform,
                location=loc,
                interview_location=loc,
                hr_call_verification=payload.hr_call_verification,
                candidate_requested_date_time=payload.candidate_requested_date_time,
                candidate_requested_date=payload.candidate_requested_date,
                candidate_requested_time=payload.candidate_requested_time,
                candidate_requested_role=payload.candidate_requested_role,
                salary_requested=payload.salary_requested,
                final_fit_salary=payload.final_fit_salary,
                joining_date=payload.joining_date,
                interview_document_files=payload.interview_document_files,
                recommendation=payload.recommendation or "Pending",
                status=InterviewStatus.SCHEDULED,
                client_rating=payload.client_rating,
                client_feedback=payload.client_feedback,
                client_strengths=payload.client_strengths,
                client_weaknesses=payload.client_weaknesses,
                client_recommendation=payload.client_recommendation,
                client_notes=payload.client_notes,
                client_id=payload.client_id,
                client_name=payload.client_name,
                client_email=payload.client_email,
                client_feedback_date=payload.client_feedback_date,
                interviewers=[i.model_dump() for i in payload.interviewers] if payload.interviewers else [],
                clients=[c.model_dump() for c in payload.clients] if payload.clients else [],
                notes=payload.notes,
                created_by=created_by,
                updated_by=created_by,
            )

            created = await self.interview_repo.create(interview_doc.to_dict())
            created_interviews.append(InterviewResponse.model_validate(created))

        logger.info(f"Batch scheduled {len(created_interviews)} interviews for job '{payload.job_title}'")
        return created_interviews

    async def get_interview_by_id(self, interview_id: str) -> InterviewResponse:
        """Fetch details for a single interview by ID."""
        interview = await self.interview_repo.get_by_id(interview_id)
        if not interview:
            raise NotFoundError("Interview not found.")
        return InterviewResponse.model_validate(interview)

    async def filter_interviews(
        self,
        candidate_id: Optional[str] = None,
        interviewer_id: Optional[str] = None,
        client_id: Optional[str] = None,
        status: Optional[Any] = None,
        interview_type: Optional[Any] = None,
        interview_type_id: Optional[str] = None,
        job_title: Optional[str] = None,
        name: Optional[str] = None,
        email: Optional[str] = None,
        scheduled_date: Optional[str] = None,
        search: Optional[str] = None,
        date_from: Optional[str] = None,
        date_to: Optional[str] = None,
        recommendation: Optional[str] = None,
        page: int = 1,
        limit: int = 10,
        skip: Optional[int] = None,
    ) -> InterviewListResponse:
        """Filter and list interviews matching parameters with pagination."""
        import math

        if skip is None:
            skip = (page - 1) * limit if page >= 1 else 0

        interviews = await self.interview_repo.filter_interviews(
            candidate_id=candidate_id,
            interviewer_id=interviewer_id,
            client_id=client_id,
            status=status,
            interview_type=interview_type,
            interview_type_id=interview_type_id,
            job_title=job_title,
            name=name,
            email=email,
            scheduled_date=scheduled_date,
            search=search,
            date_from=date_from,
            date_to=date_to,
            recommendation=recommendation,
            skip=skip,
            limit=limit,
        )
        total = await self.interview_repo.count_interviews(
            candidate_id=candidate_id,
            interviewer_id=interviewer_id,
            client_id=client_id,
            status=status,
            interview_type=interview_type,
            interview_type_id=interview_type_id,
            job_title=job_title,
            name=name,
            email=email,
            scheduled_date=scheduled_date,
            search=search,
            date_from=date_from,
            date_to=date_to,
            recommendation=recommendation,
        )
        items = [InterviewResponse.model_validate(i) for i in interviews]
        total_pages = math.ceil(total / limit) if limit > 0 else 1

        return InterviewListResponse(
            total=total,
            interviews=items,
            page=page,
            limit=limit,
            total_pages=total_pages,
        )

    async def update_interview(
        self,
        interview_id: str,
        payload: InterviewUpdateRequest,
        updated_by: Optional[str] = None,
    ) -> InterviewResponse:
        """Update fields of an existing interview document."""
        existing = await self.interview_repo.get_by_id(interview_id)
        if not existing:
            raise NotFoundError("Interview not found.")

        update_fields = payload.model_dump(exclude_unset=True)
        if "status" in update_fields and hasattr(update_fields["status"], "value"):
            update_fields["status"] = update_fields["status"].value
        if "interview_type" in update_fields and hasattr(update_fields["interview_type"], "value"):
            update_fields["interview_type"] = update_fields["interview_type"].value

        check_date = update_fields.get("scheduled_date") or existing.get("scheduled_date")
        check_time = update_fields.get("scheduled_time") or existing.get("scheduled_time")
        if check_date and check_time:
            conflict_res = await self.interview_repo.check_schedule_conflict(
                scheduled_date=check_date,
                scheduled_time=check_time,
                interviewer_id=update_fields.get("interviewer_id") or existing.get("interviewer_id"),
                interviewer_name=update_fields.get("interviewer_name") or existing.get("interviewer_name"),
                client_id=update_fields.get("client_id") or existing.get("client_id"),
                client_name=update_fields.get("client_name") or existing.get("client_name"),
                interviewers=update_fields.get("interviewers") if "interviewers" in update_fields else existing.get("interviewers"),
                clients=update_fields.get("clients") if "clients" in update_fields else existing.get("clients"),
                exclude_interview_id=interview_id,
            )
            if conflict_res.get("has_conflict"):
                raise HTTPException(status_code=400, detail=conflict_res.get("conflict_message"))

        update_fields["updated_by"] = updated_by
        update_fields["updated_at"] = utc_now().isoformat()

        updated_doc = await self.interview_repo.update(interview_id, update_fields)
        logger.info(f"Updated interview record '{interview_id}'")
        return InterviewResponse.model_validate(updated_doc)

    async def reschedule_interview(
        self,
        interview_id: str,
        payload: InterviewRescheduleRequest,
        updated_by: Optional[str] = None,
    ) -> InterviewResponse:
        """Reschedule an existing interview and push entry into reschedule history."""
        existing = await self.interview_repo.get_by_id(interview_id)
        if not existing:
            raise NotFoundError("Interview not found.")

        # Validation: Check for interviewer/client schedule time conflict
        conflict_res = await self.interview_repo.check_schedule_conflict(
            scheduled_date=payload.scheduled_date,
            scheduled_time=payload.scheduled_time,
            interviewer_id=existing.get("interviewer_id"),
            interviewer_name=existing.get("interviewer_name"),
            client_id=existing.get("client_id"),
            client_name=existing.get("client_name"),
            interviewers=existing.get("interviewers"),
            clients=existing.get("clients"),
            exclude_interview_id=interview_id,
        )
        if conflict_res.get("has_conflict"):
            raise HTTPException(status_code=400, detail=conflict_res.get("conflict_message"))

        history_entry = {
            "previous_date": existing.get("scheduled_date"),
            "previous_time": existing.get("scheduled_time"),
            "new_date": payload.scheduled_date,
            "new_time": payload.scheduled_time,
            "reason": payload.reason,
            "rescheduled_at": utc_now().isoformat(),
            "rescheduled_by": updated_by,
        }

        updated_doc = await self.interview_repo.reschedule_interview(
            interview_id=interview_id,
            new_date=payload.scheduled_date,
            new_time=payload.scheduled_time,
            timezone=payload.timezone or existing.get("timezone", "Asia/Kolkata"),
            duration_minutes=payload.duration_minutes or existing.get("duration_minutes", 60),
            history_entry=history_entry,
            updated_by=updated_by,
        )

        logger.info(f"Rescheduled interview ID '{interview_id}' to date {payload.scheduled_date} at {payload.scheduled_time}")
        return InterviewResponse.model_validate(updated_doc)

    async def submit_feedback(
        self,
        interview_id: str,
        payload: InterviewFeedbackRequest,
        updated_by: Optional[str] = None,
    ) -> InterviewResponse:
        """Submit rating, feedback, strengths, and weaknesses for an interview (Round Interviewer and/or Client)."""
        existing = await self.interview_repo.get_by_id(interview_id)
        if not existing:
            raise NotFoundError("Interview not found.")

        updated_doc = await self.interview_repo.submit_feedback(
            interview_id=interview_id,
            rating=payload.rating,
            feedback=payload.feedback,
            strengths=payload.strengths,
            weaknesses=payload.weaknesses,
            recommendation=payload.recommendation,
            notes=payload.notes,
            client_rating=payload.client_rating,
            client_feedback=payload.client_feedback,
            client_strengths=payload.client_strengths,
            client_weaknesses=payload.client_weaknesses,
            client_recommendation=payload.client_recommendation,
            client_notes=payload.client_notes,
            client_name=payload.client_name,
            client_feedback_date=payload.client_feedback_date,
            updated_by=updated_by,
            candidate_requested_date_time=payload.candidate_requested_date_time,
            candidate_requested_date=payload.candidate_requested_date,
            candidate_requested_time=payload.candidate_requested_time,
            candidate_requested_role=payload.candidate_requested_role,
            salary_requested=payload.salary_requested,
            final_fit_salary=payload.final_fit_salary,
            joining_date=payload.joining_date,
            interview_document_files=payload.interview_document_files,
            interview_feedback_files=payload.interview_feedback_files,
            interviewers=[i.model_dump() for i in payload.interviewers] if payload.interviewers else None,
            clients=[c.model_dump() for c in payload.clients] if payload.clients else None,
            skill_ratings=[s.model_dump() for s in payload.skill_ratings] if payload.skill_ratings else None,
            category_scores=[cat.model_dump() for cat in payload.category_scores] if payload.category_scores else None,
            ai_score=payload.ai_score,
            ai_recommendation=payload.ai_recommendation,
            hr_call_verification=payload.hr_call_verification,
            location=payload.location,
            interview_location=payload.interview_location,
            meeting_link=payload.meeting_link,
            meeting_platform=payload.meeting_platform,
            evaluator_user_id=payload.evaluator_user_id or updated_by,
            evaluator_name=payload.evaluator_name,
            evaluator_email=payload.evaluator_email,
            evaluator_role=payload.evaluator_role,
        )

        logger.info(f"Submitted feedback for interview ID '{interview_id}'")
        return InterviewResponse.model_validate(updated_doc)

    async def get_next_round_number(
        self, candidate_id: str, interview_type: Optional[str] = None, interview_type_id: Optional[str] = None
    ) -> Dict[str, Any]:
        """Calculate the next round number for a candidate, specifically per interview_type or interview_type_id if provided."""
        raw_interviews = await self.interview_repo.get_by_candidate_id(candidate_id, skip=0, limit=200)
        
        if not raw_interviews:
            return {"candidate_id": candidate_id, "interview_type": interview_type, "interview_type_id": interview_type_id, "next_round_number": 1}

        matching_rounds = []
        if interview_type_id or interview_type:
            target_id = (interview_type_id or "").strip()
            norm_type = (interview_type or "").upper().strip()
            for doc in raw_interviews:
                doc_type_id = str(doc.get("interview_type_id", "")).strip()
                doc_type = str(doc.get("interview_type", "")).upper().strip()
                if (target_id and doc_type_id == target_id) or (norm_type and doc_type == norm_type):
                    matching_rounds.append(doc.get("round_number", 1))
        
        if matching_rounds:
            next_round = max(matching_rounds) + 1
        else:
            next_round = max((doc.get("round_number", 1) for doc in raw_interviews), default=0) + 1 if not (interview_type or interview_type_id) else 1

        return {
            "candidate_id": candidate_id,
            "interview_type": interview_type,
            "interview_type_id": interview_type_id,
            "next_round_number": next_round,
            "total_existing_rounds": len(raw_interviews)
        }

    async def bulk_submit_feedback(
        self,
        payload: BulkInterviewFeedbackRequest,
        updated_by: Optional[str] = None,
    ) -> List[InterviewResponse]:
        """Submit feedback for multiple candidate interviews in bulk."""
        updated_interviews: List[InterviewResponse] = []

        for item in payload.items:
            res = await self.submit_feedback(
                interview_id=item.interview_id,
                payload=item,
                updated_by=updated_by,
            )
            updated_interviews.append(res)

        logger.info(f"Bulk submitted feedback for {len(updated_interviews)} interviews.")
        return updated_interviews



    async def delete_interview(self, interview_id: str) -> bool:
        """Delete interview record from database."""
        existing = await self.interview_repo.get_by_id(interview_id)
        if not existing:
            raise NotFoundError("Interview not found.")

        deleted = await self.interview_repo.delete(interview_id)
        logger.info(f"Deleted interview document ID '{interview_id}'")
        return deleted

    async def get_candidate_history(self, candidate_id: str) -> CandidateFullHistoryResponse:
        """Retrieve complete candidate profile and all interview rounds sorted by round number."""
        raw_interviews = await self.interview_repo.get_by_candidate_id(candidate_id, skip=0, limit=200)
        if not raw_interviews:
            raise NotFoundError(f"No interview records found for candidate ID '{candidate_id}'.")

        # Sort rounds by round_number ascending
        sorted_docs = sorted(raw_interviews, key=lambda x: x.get("round_number", 1))
        
        # Aggregate candidate info across all rounds to make sure we don't display empty/null values at top level
        candidate_name = next((d.get("candidate_name") for d in sorted_docs if d.get("candidate_name")), "Unknown")
        candidate_email = next((d.get("candidate_email") for d in sorted_docs if d.get("candidate_email")), None)
        resume_id = next((d.get("resume_id") for d in sorted_docs if d.get("resume_id")), None)
        job_id = next((d.get("job_id") for d in sorted_docs if d.get("job_id")), None)
        job_title = next((d.get("job_title") for d in sorted_docs if d.get("job_title")), None)
        job_location = next((d.get("job_location") for d in sorted_docs if d.get("job_location")), None)
        job_type = next((d.get("job_type") for d in sorted_docs if d.get("job_type")), None)
        
        location = next((d.get("location") for d in sorted_docs if d.get("location")), None)
        interview_location = next((d.get("interview_location") for d in sorted_docs if d.get("interview_location")), None)
        
        hr_call_verification = next((d.get("hr_call_verification") for d in sorted_docs if d.get("hr_call_verification") and d.get("hr_call_verification") != "Pending"), "Pending")
        if hr_call_verification == "Pending":
            hr_call_verification = next((d.get("hr_call_verification") for d in sorted_docs if d.get("hr_call_verification")), "Pending")

        candidate_requested_date = next((d.get("candidate_requested_date") for d in sorted_docs if d.get("candidate_requested_date")), None)
        candidate_requested_time = next((d.get("candidate_requested_time") for d in sorted_docs if d.get("candidate_requested_time")), None)
        candidate_requested_role = next((d.get("candidate_requested_role") for d in sorted_docs if d.get("candidate_requested_role")), None)
        salary_requested = next((d.get("salary_requested") for d in sorted_docs if d.get("salary_requested")), None)
        final_fit_salary = next((d.get("final_fit_salary") for d in sorted_docs if d.get("final_fit_salary")), None)
        joining_date = next((d.get("joining_date") for d in sorted_docs if d.get("joining_date")), None)

        merged_files = []
        merged_feedback_files = []
        for d in sorted_docs:
            files = d.get("interview_document_files") or []
            for f in files:
                if f not in merged_files:
                    merged_files.append(f)
            fb_files = d.get("interview_feedback_files") or []
            for fb in fb_files:
                if fb not in merged_feedback_files:
                    merged_feedback_files.append(fb)

        rounds = [InterviewResponse.model_validate(doc) for doc in sorted_docs]

        return CandidateFullHistoryResponse(
            candidate_id=candidate_id,
            candidate_name=candidate_name,
            candidate_email=candidate_email,
            resume_id=resume_id,
            job_id=job_id,
            job_title=job_title,
            job_location=job_location,
            job_type=job_type,
            location=location or interview_location,
            interview_location=interview_location or location,
            hr_call_verification=hr_call_verification,
            candidate_requested_date=candidate_requested_date,
            candidate_requested_time=candidate_requested_time,
            candidate_requested_role=candidate_requested_role,
            salary_requested=salary_requested,
            final_fit_salary=final_fit_salary,
            joining_date=joining_date,
            interview_document_files=merged_files,
            interview_feedback_files=merged_feedback_files,
            total_rounds=len(rounds),
            rounds=rounds,
        )

    async def send_interview_email(
        self,
        interview_id: str,
        payload: SendInterviewEmailRequest,
    ) -> Dict[str, Any]:
        """Send interview schedule notification emails to candidate and/or interviewer using dynamic template values."""
        existing = await self.interview_repo.get_by_id(interview_id)
        if not existing:
            raise NotFoundError("Interview not found.")

        template_service = TemplateService()
        await template_service.initialize_default_templates()

        # Determine emails
        candidate_email = payload.candidate_email or existing.get("candidate_email")
        if not candidate_email and existing.get("resume_id"):
            try:
                db = self.interview_repo.collection.database
                resume_doc = await db.resumes.find_one({"_id": existing.get("resume_id")})
                if resume_doc:
                    parsed = resume_doc.get("parsed_data") or {}
                    candidate_email = parsed.get("email") or resume_doc.get("email")
            except Exception as err:
                logger.warning(f"Could not fetch resume email: {err}")

        interviewer_email = payload.interviewer_email or existing.get("interviewer_email")

        # Fetch SMTP config
        settings_service = SettingsService()
        config_model = await settings_service.repository.get_email_config()
        if not config_model:
            raise HTTPException(status_code=400, detail="Email configuration not set. Please configure SMTP settings in Settings first.")

        # Prepare replacement dictionary
        notes_content = payload.custom_notes or existing.get("notes") or "N/A"
        meeting_link = existing.get("meeting_link") or existing.get("location") or "Will be shared shortly"

        doc_files = existing.get("interview_document_files") or []
        fb_files = existing.get("interview_feedback_files") or []
        combined_attachment_urls = list(dict.fromkeys([f for f in (doc_files + fb_files) if f and str(f).strip()]))

        replacements = {
            "candidate_name": existing.get("candidate_name", ""),
            "candidate_email": candidate_email or "N/A",
            "interviewer_name": existing.get("interviewer_name", "Interviewer"),
            "interviewer_email": interviewer_email or "N/A",
            "job_title": existing.get("job_title", ""),
            "scheduled_date": existing.get("scheduled_date", ""),
            "scheduled_time": existing.get("scheduled_time", ""),
            "timezone": existing.get("timezone", "Asia/Kolkata"),
            "duration_minutes": str(existing.get("duration_minutes", 60)),
            "meeting_platform": existing.get("meeting_platform", "Google Meet"),
            "meeting_link": meeting_link,
            "location": existing.get("location") or existing.get("interview_location") or "Online",
            "interview_type": str(existing.get("interview_type", "")).replace("_", " "),
            "round_number": str(existing.get("round_number", 1)),
            "notes": notes_content,
            "company_name": config_model.sender_name or "HR Team",
        }

        def render_text(text: str) -> str:
            rendered = text
            for key, val in replacements.items():
                rendered = rendered.replace(f"{{{{{key}}}}}", str(val))
            return rendered

        all_templates = await template_service.get_all_templates()
        sent_recipients = []

        try:
            password = decrypt_password(config_model.smtp_password)
            
            # Helper function for executive HTML layout
            def build_professional_email_html(
                header_badge: str,
                header_title: str,
                body_content: str,
                schedule_items: List[Dict[str, str]],
                meeting_url: str = "",
                notes_text: Optional[str] = None,
                attached_files: Optional[List[str]] = None,
                sender_company: str = "Recruitment Team"
            ) -> str:
                rows_html = ""
                for item in schedule_items:
                    label = item.get("label", "")
                    val = item.get("value", "")
                    if val and val != "N/A":
                        rows_html += f"""
                        <tr>
                          <td style="padding: 10px 12px; color: #64748b; font-weight: 600; width: 140px; border-bottom: 1px solid #f1f5f9; vertical-align: top;">{label}</td>
                          <td style="padding: 10px 12px; color: #0f172a; font-weight: 700; border-bottom: 1px solid #f1f5f9; vertical-align: top;">{val}</td>
                        </tr>
                        """

                cta_html = ""
                if meeting_url and meeting_url.startswith("http"):
                    cta_html = f"""
                    <tr>
                      <td align="center" style="padding: 0 36px 28px 36px;">
                        <a href="{meeting_url}" target="_blank" style="display: inline-block; background: linear-gradient(135deg, #4f46e5 0%, #4338ca 100%); color: #ffffff; font-size: 14px; font-weight: 700; text-decoration: none; padding: 14px 32px; border-radius: 10px; box-shadow: 0 4px 12px rgba(79, 70, 229, 0.3);">
                          🎥 Join Interview Meeting &rarr;
                        </a>
                        <div style="font-size: 11px; color: #94a3b8; margin-top: 10px;">Direct Link: <a href="{meeting_url}" style="color: #4f46e5; text-decoration: underline;">{meeting_url}</a></div>
                      </td>
                    </tr>
                    """
                elif meeting_url and meeting_url != "N/A" and meeting_url != "Will be shared shortly":
                    rows_html += f"""
                    <tr>
                      <td style="padding: 10px 12px; color: #64748b; font-weight: 600; width: 140px; border-bottom: 1px solid #f1f5f9; vertical-align: top;">Location</td>
                      <td style="padding: 10px 12px; color: #0f172a; font-weight: 700; border-bottom: 1px solid #f1f5f9; vertical-align: top;">{meeting_url}</td>
                    </tr>
                    """

                files_html = ""
                if attached_files and len(attached_files) > 0:
                    file_badges = ""
                    for file_url in attached_files:
                        if not file_url or not file_url.strip():
                            continue
                        f_str = file_url.strip()
                        raw_name = f_str.split("/")[-1] if "/" in f_str else f_str
                        disp_name = urllib.parse.unquote(raw_name)
                        disp_name = re.sub(r'^[a-f0-9]{8,32}_', '', disp_name)

                        if f_str.startswith("http://") or f_str.startswith("https://"):
                            file_badges += f"""
                            <a href="{f_str}" target="_blank" style="display: inline-block; background-color: #ffffff; border: 1px solid #cbd5e1; color: #4f46e5; font-weight: 700; font-size: 12px; text-decoration: none; padding: 7px 14px; border-radius: 8px; margin: 4px 6px 4px 0; box-shadow: 0 1px 2px rgba(0,0,0,0.05);">
                              📄 {disp_name} &rarr;
                            </a>
                            """
                        else:
                            file_badges += f"""
                            <span style="display: inline-block; background-color: #f1f5f9; border: 1px solid #cbd5e1; color: #334155; font-weight: 600; font-size: 12px; padding: 7px 14px; border-radius: 8px; margin: 4px 6px 4px 0;">
                              📄 {disp_name}
                            </span>
                            """

                    if file_badges:
                        files_html = f"""
                        <tr>
                          <td style="padding: 0 36px 24px 36px;">
                            <div style="background-color: #f8fafc; border: 1px solid #e2e8f0; border-left: 4px solid #3b82f6; border-radius: 10px; padding: 16px 20px; font-size: 13px; color: #0f172a;">
                              <div style="font-weight: 800; font-size: 11px; text-transform: uppercase; letter-spacing: 0.8px; margin-bottom: 8px; color: #2563eb;">
                                📎 Attached Session Files & S3 Resources ({len(attached_files)})
                              </div>
                              <div>
                                {file_badges}
                              </div>
                            </div>
                          </td>
                        </tr>
                        """

                notes_html = ""
                if notes_text and notes_text.strip() and notes_text.strip() != "N/A":
                    notes_html = f"""
                    <tr>
                      <td style="padding: 0 36px 28px 36px;">
                        <div style="background-color: #fffbeb; border: 1px solid #fef3c7; border-left: 4px solid #f59e0b; border-radius: 10px; padding: 16px 20px; font-size: 13px; color: #92400e;">
                          <div style="font-weight: 800; font-size: 11px; text-transform: uppercase; letter-spacing: 0.8px; margin-bottom: 6px; color: #b45309;">📌 Schedule Notes</div>
                          <div style="line-height: 1.5; color: #78350f; font-weight: 500;">{notes_text.strip()}</div>
                        </div>
                      </td>
                    </tr>
                    """

                return f"""<!DOCTYPE html>
<html>
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
</head>
<body style="margin: 0; padding: 0; background-color: #f1f5f9; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, 'Helvetica Neue', Arial, sans-serif;">
  <table role="presentation" width="100%" border="0" cellspacing="0" cellpadding="0" style="background-color: #f1f5f9; padding: 32px 12px;">
    <tr>
      <td align="center">
        <table role="presentation" width="100%" border="0" cellspacing="0" cellpadding="0" style="max-width: 580px; background-color: #ffffff; border-radius: 16px; overflow: hidden; border: 1px solid #e2e8f0; box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.05);">
          
          <!-- Header Banner -->
          <tr>
            <td style="background: linear-gradient(135deg, #4f46e5 0%, #3b82f6 100%); padding: 28px 36px; text-align: left;">
              <span style="display: inline-block; background-color: rgba(255, 255, 255, 0.2); color: #ffffff; font-size: 10px; font-weight: 800; text-transform: uppercase; letter-spacing: 1.2px; padding: 4px 12px; border-radius: 20px;">{header_badge}</span>
              <h1 style="color: #ffffff; font-size: 20px; font-weight: 800; margin: 10px 0 0 0; text-shadow: 0 1px 2px rgba(0,0,0,0.1);">{header_title}</h1>
            </td>
          </tr>

          <!-- Main Text Body -->
          <tr>
            <td style="padding: 32px 36px 20px 36px; color: #334155; font-size: 14px; line-height: 1.6;">
              {body_content}
            </td>
          </tr>

          <!-- Schedule Table Card -->
          <tr>
            <td style="padding: 0 36px 24px 36px;">
              <table role="presentation" width="100%" border="0" cellspacing="0" cellpadding="0" style="background-color: #f8fafc; border-radius: 12px; border: 1px solid #e2e8f0; border-left: 4px solid #4f46e5;">
                <tr>
                  <td style="padding: 18px 20px;">
                    <div style="font-size: 11px; font-weight: 800; color: #4f46e5; text-transform: uppercase; letter-spacing: 0.8px; margin-bottom: 10px;">📅 Interview Details</div>
                    <table role="presentation" width="100%" border="0" cellspacing="0" cellpadding="0" style="font-size: 13px; border-collapse: collapse;">
                      {rows_html}
                    </table>
                  </td>
                </tr>
              </table>
            </td>
          </tr>

          <!-- CTA Button -->
          {cta_html}

          <!-- Attached Session Files -->
          {files_html}

          <!-- Schedule Notes -->
          {notes_html}

          <!-- Footer Signature -->
          <tr>
            <td style="background-color: #f8fafc; border-top: 1px solid #e2e8f0; padding: 20px 36px; text-align: center; color: #64748b; font-size: 12px;">
              <p style="margin: 0; font-weight: 700; color: #334155;">{sender_company}</p>
              <p style="margin: 4px 0 0 0; color: #94a3b8; font-size: 11px;">Automated Interview System • AI Recruiter</p>
            </td>
          </tr>

        </table>
      </td>
    </tr>
  </table>
</body>
</html>"""

            # 1. SEND TO CANDIDATE
            if payload.send_to_candidate:
                if not candidate_email:
                    raise HTTPException(status_code=400, detail="Candidate email is missing. Please enter candidate email.")

                cand_addrs = [a.strip() for a in candidate_email.replace(";", ",").split(",") if a.strip()]
                if not cand_addrs:
                    raise HTTPException(status_code=400, detail="Candidate email is invalid or empty.")

                cand_template = None
                if payload.template_id:
                    cand_template = await template_service.get_template(payload.template_id)
                if not cand_template:
                    cand_template = next((t for t in all_templates if "Invitation" in t.name or "Candidate" in t.name), all_templates[0] if all_templates else None)

                cand_subject = render_text(cand_template.subject if cand_template else f"Interview Invitation: {existing.get('job_title')} - {existing.get('candidate_name')}")
                cand_raw_body = render_text(cand_template.body if cand_template else f"<p>Dear <b>{existing.get('candidate_name')}</b>,</p><p>We are pleased to invite you for an interview for the <b>{existing.get('job_title')}</b> position.</p>")

                cand_schedule_items = [
                    {"label": "Candidate", "value": existing.get("candidate_name", "")},
                    {"label": "Job Position", "value": existing.get("job_title", "")},
                    {"label": "Interview Round", "value": f"Round {existing.get('round_number', 1)} ({str(existing.get('interview_type', '')).replace('_', ' ')})"},
                    {"label": "Scheduled Date", "value": replacements['scheduled_date']},
                    {"label": "Scheduled Time", "value": f"{replacements['scheduled_time']} ({replacements['timezone']})"},
                    {"label": "Duration", "value": f"{replacements['duration_minutes']} Minutes"},
                    {"label": "Interviewer", "value": replacements['interviewer_name']},
                ]

                cand_final_html = build_professional_email_html(
                    header_badge="Candidate Invitation",
                    header_title=f"Interview Invitation - {existing.get('job_title')}",
                    body_content=cand_raw_body,
                    schedule_items=cand_schedule_items,
                    meeting_url=meeting_link,
                    notes_text=notes_content,
                    attached_files=combined_attachment_urls,
                    sender_company=config_model.sender_name or "Recruitment Team"
                )

                msg_cand = EmailMessage()
                msg_cand["Subject"] = cand_subject
                msg_cand["From"] = f"{config_model.sender_name} <{config_model.sender_email}>"
                msg_cand["To"] = ", ".join(cand_addrs)
                msg_cand.set_content("Please enable HTML to view this email.")
                msg_cand.add_alternative(cand_final_html, subtype='html')

                if config_model.use_ssl:
                    server = smtplib.SMTP_SSL(config_model.smtp_server, config_model.smtp_port)
                else:
                    server = smtplib.SMTP(config_model.smtp_server, config_model.smtp_port)
                    if config_model.use_tls:
                        server.starttls()

                server.login(config_model.smtp_username, password)
                server.send_message(msg_cand, to_addrs=cand_addrs)
                server.quit()
                sent_recipients.append(f"Candidate ({', '.join(cand_addrs)})")

            # 2. SEND TO INTERVIEWER
            if payload.send_to_interviewer:
                if not interviewer_email:
                    raise HTTPException(status_code=400, detail="Interviewer email is missing. Please enter interviewer email.")

                int_addrs = [a.strip() for a in interviewer_email.replace(";", ",").split(",") if a.strip()]
                if not int_addrs:
                    raise HTTPException(status_code=400, detail="Interviewer email is invalid or empty.")

                interviewer_template = next((t for t in all_templates if "Interviewer" in t.name), None)

                if interviewer_template:
                    int_subject = render_text(interviewer_template.subject)
                    int_raw_body = render_text(interviewer_template.body)
                else:
                    int_subject = f"Interview Assigned: {existing.get('job_title')} - {existing.get('candidate_name')}"
                    int_raw_body = f"""
                    <p>Hello <b>{replacements['interviewer_name']}</b>,</p>
                    <p>You have been assigned to conduct an interview with candidate <b>{replacements['candidate_name']}</b> for the <b>{replacements['job_title']}</b> position.</p>
                    <p>Please ensure to update candidate rating and feedback post-interview.</p>
                    """

                int_schedule_items = [
                    {"label": "Candidate", "value": f"{existing.get('candidate_name', '')} ({candidate_email or 'N/A'})"},
                    {"label": "Job Position", "value": existing.get("job_title", "")},
                    {"label": "Interview Round", "value": f"Round {existing.get('round_number', 1)} ({str(existing.get('interview_type', '')).replace('_', ' ')})"},
                    {"label": "Scheduled Date", "value": replacements['scheduled_date']},
                    {"label": "Scheduled Time", "value": f"{replacements['scheduled_time']} ({replacements['timezone']})"},
                    {"label": "Duration", "value": f"{replacements['duration_minutes']} Minutes"},
                ]

                int_final_html = build_professional_email_html(
                    header_badge="Interviewer Assignment",
                    header_title=f"Interview Assigned: {existing.get('candidate_name')}",
                    body_content=int_raw_body,
                    schedule_items=int_schedule_items,
                    meeting_url=meeting_link,
                    notes_text=notes_content,
                    attached_files=combined_attachment_urls,
                    sender_company=config_model.sender_name or "Recruitment Team"
                )

                msg_int = EmailMessage()
                msg_int["Subject"] = int_subject
                msg_int["From"] = f"{config_model.sender_name} <{config_model.sender_email}>"
                msg_int["To"] = ", ".join(int_addrs)
                msg_int.set_content("Please enable HTML to view this email.")
                msg_int.add_alternative(int_final_html, subtype='html')

                if config_model.use_ssl:
                    server = smtplib.SMTP_SSL(config_model.smtp_server, config_model.smtp_port)
                else:
                    server = smtplib.SMTP(config_model.smtp_server, config_model.smtp_port)
                    if config_model.use_tls:
                        server.starttls()

                server.login(config_model.smtp_username, password)
                server.send_message(msg_int, to_addrs=int_addrs)
                server.quit()
                sent_recipients.append(f"Interviewer ({', '.join(int_addrs)})")

            if not sent_recipients:
                raise HTTPException(status_code=400, detail="No email recipients selected.")

            # Record email dispatch in database (updates email_sent_count, last_email_sent_at, email_sent_history)
            email_type = "RESCHEDULE_NOTIFICATION" if "RESCHEDULED" in (payload.custom_notes or "").upper() else "SCHEDULE_NOTIFICATION"
            await self.interview_repo.record_email_sent(
                interview_id=interview_id,
                sent_recipients=sent_recipients,
                email_type=email_type,
            )

        except HTTPException:
            raise
        except Exception as e:
            logger.error(f"Failed to send interview email: {e}")
            raise HTTPException(status_code=500, detail=f"Failed to send email: {str(e)}")

        logger.info(f"Sent interview email for interview '{interview_id}' to: {sent_recipients}")
        return {
            "status": "success",
            "message": f"Interview email sent successfully to {', '.join(sent_recipients)}",
            "recipients": sent_recipients,
        }

    async def share_interview_summary(self, payload: ShareInterviewSummaryRequest) -> Dict[str, Any]:
        """Email a summary table (candidate details, round, status, rating, outcome, feedback) of selected sessions."""
        ids = list(dict.fromkeys(i for i in payload.interview_ids if i and i.strip()))
        docs = await self.interview_repo.find_many(query={"id": {"$in": ids}}, limit=len(ids))
        if not docs:
            raise NotFoundError("No interviews found for the selected sessions.")
        order = {interview_id: idx for idx, interview_id in enumerate(ids)}
        docs.sort(key=lambda d: order.get(str(d.get("id")), len(order)))

        settings_service = SettingsService()
        config_model = await settings_service.repository.get_email_config()
        if not config_model:
            raise HTTPException(status_code=400, detail="Email configuration not set. Please configure SMTP settings in Settings first.")

        # Resolve interview type labels from configured types (by id, then code)
        type_labels: Dict[str, str] = {}
        try:
            db = self.interview_repo.collection.database
            async for t in db.interview_types.find({}, {"id": 1, "code": 1, "name": 1}):
                if t.get("id"):
                    type_labels[str(t["id"])] = t.get("name", "")
                if t.get("code"):
                    type_labels[str(t["code"]).upper()] = t.get("name", "")
        except Exception as err:
            logger.warning(f"Could not load interview types for summary email: {err}")

        today_str = utc_now().date().isoformat()

        def esc(value: Any) -> str:
            return html.escape(str(value)) if value not in (None, "") else "-"

        def session_status(doc: Dict[str, Any]) -> str:
            # Mirrors the Completed & Past Session History card badges on the Interview Dashboard
            rec = (doc.get("recommendation") or doc.get("client_recommendation") or doc.get("ai_recommendation") or "").strip()
            if doc.get("status") == InterviewStatus.COMPLETED.value:
                return "Completed"
            if "reschedule" in rec.lower():
                return "Reschedule Required"
            if doc.get("scheduled_date") and doc["scheduled_date"] < today_str:
                return "Past / Missed"
            return str(doc.get("status") or "-").replace("_", " ").title()

        status_colors = {
            "Completed": ("#d1fae5", "#065f46"),
            "Reschedule Required": ("#ffe4e6", "#9f1239"),
            "Past / Missed": ("#fef3c7", "#92400e"),
        }

        cell = "padding: 8px 10px; border-bottom: 1px solid #e2e8f0; vertical-align: top; font-size: 12px; color: #0f172a;"
        rows_html = ""
        export_rows: List[List[Any]] = []
        for doc in docs:
            raw_type = str(doc.get("interview_type") or "")
            round_label = (
                type_labels.get(str(doc.get("interview_type_id") or ""))
                or type_labels.get(raw_type.upper())
                or raw_type.replace("_", " ")
                or "-"
            )
            status_label = session_status(doc)
            bg, fg = status_colors.get(status_label, ("#f1f5f9", "#334155"))
            rating = doc.get("rating") or doc.get("client_rating")
            outcome = doc.get("recommendation") or doc.get("client_recommendation") or doc.get("ai_recommendation")
            feedback_parts = []
            if doc.get("feedback"):
                feedback_parts.append(f"<b>Interviewer:</b> {html.escape(str(doc['feedback']))}")
            if doc.get("client_feedback"):
                feedback_parts.append(f"<b>Client:</b> {html.escape(str(doc['client_feedback']))}")
            feedback_html = "<br><br>".join(feedback_parts).replace("\n", "<br>") or "-"
            export_rows.append([
                doc.get("candidate_name"), doc.get("candidate_email"), doc.get("job_title"), round_label,
                doc.get("round_number", 1), doc.get("scheduled_date"), doc.get("scheduled_time"),
                doc.get("interviewer_name") or doc.get("client_name"), status_label, rating, outcome,
                doc.get("feedback"), doc.get("client_feedback"),
            ])

            rows_html += f"""
            <tr>
              <td style="{cell}"><b>{esc(doc.get("candidate_name"))}</b><br><span style="color:#64748b;">{esc(doc.get("candidate_email"))}</span></td>
              <td style="{cell}">{esc(doc.get("job_title"))}</td>
              <td style="{cell}">{esc(round_label)} (R{esc(doc.get("round_number", 1))})</td>
              <td style="{cell} white-space: nowrap;">{esc(doc.get("scheduled_date"))}<br><span style="color:#64748b;">{esc(doc.get("scheduled_time"))}</span></td>
              <td style="{cell}">{esc(doc.get("interviewer_name") or doc.get("client_name"))}</td>
              <td style="{cell}"><span style="display:inline-block; background:{bg}; color:{fg}; font-weight:700; padding:2px 8px; border-radius:10px; white-space:nowrap;">{esc(status_label)}</span></td>
              <td style="{cell} white-space: nowrap;">{f"{esc(rating)} / 5" if rating else "-"}</td>
              <td style="{cell}">{esc(outcome)}</td>
              <td width="30%" style="{cell} color:#334155; min-width:260px;">{feedback_html}</td>
            </tr>
            """

        header_cell = "padding: 8px 10px; background: #f1f5f9; color: #475569; font-size: 11px; text-transform: uppercase; letter-spacing: 0.5px; text-align: left; border-bottom: 1px solid #cbd5e1;"
        headers = ["Candidate", "Job Title", "Round", "Scheduled", "Interviewer", "Status", "Rating", "Outcome", "Feedback Comments"]
        # Feedback text is the longest column; give it a fixed share so it doesn't get squeezed
        header_html = "".join(
            f'<th width="30%" style="{header_cell} min-width: 260px;">{h}</th>'
            if h == "Feedback Comments"
            else f'<th style="{header_cell} white-space: nowrap;">{h}</th>'
            for h in headers
        )
        sender_company = config_model.sender_name or "Recruitment Team"
        notes_html = ""
        if payload.notes and payload.notes.strip():
            notes_html = f"""
            <div style="background:#fffbeb; border:1px solid #fef3c7; border-left:4px solid #f59e0b; border-radius:10px; padding:12px 16px; margin: 0 0 16px 0; font-size:13px; color:#78350f;">
              {html.escape(payload.notes.strip()).replace(chr(10), "<br>")}
            </div>
            """

        email_html = f"""<!DOCTYPE html>
<html>
<head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1.0"></head>
<body style="margin:0; padding:24px 12px; background:#f1f5f9; font-family:-apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Arial, sans-serif;">
  <div style="max-width:1000px; margin:0 auto; background:#ffffff; border:1px solid #e2e8f0; border-radius:16px; overflow:hidden;">
    <div style="background:linear-gradient(135deg, #4f46e5 0%, #3b82f6 100%); padding:24px 28px;">
      <span style="display:inline-block; background:rgba(255,255,255,0.2); color:#ffffff; font-size:10px; font-weight:800; text-transform:uppercase; letter-spacing:1.2px; padding:4px 12px; border-radius:20px;">Interview Status Summary</span>
      <h1 style="color:#ffffff; font-size:20px; font-weight:800; margin:10px 0 0 0;">Candidate Interview Status ({len(docs)})</h1>
    </div>
    <div style="padding:24px 28px;">
      {notes_html}
      <div style="overflow-x:auto;">
        <table role="presentation" width="100%" cellspacing="0" cellpadding="0" style="border-collapse:collapse; border:1px solid #e2e8f0;">
          <thead><tr>{header_html}</tr></thead>
          <tbody>{rows_html}</tbody>
        </table>
      </div>
    </div>
    <div style="background:#f8fafc; border-top:1px solid #e2e8f0; padding:16px 28px; text-align:center; color:#64748b; font-size:12px;">
      <p style="margin:0; font-weight:700; color:#334155;">{html.escape(sender_company)}</p>
      <p style="margin:4px 0 0 0; color:#94a3b8; font-size:11px;">Automated Interview System • AI Recruiter</p>
    </div>
  </div>
</body>
</html>"""

        to_addrs = list(dict.fromkeys(str(e) for e in payload.to_emails))
        msg = EmailMessage()
        msg["Subject"] = f"Interview Status Summary - {len(docs)} Candidate Session{'s' if len(docs) != 1 else ''}"
        msg["From"] = f"{config_model.sender_name} <{config_model.sender_email}>"
        msg["To"] = ", ".join(to_addrs)
        msg.set_content("Please enable HTML to view this email.")
        msg.add_alternative(email_html, subtype="html")

        if payload.attachment_format != "none":
            export_headers = [
                "Candidate Name", "Candidate Email", "Job Title", "Round", "Round Number", "Scheduled Date",
                "Scheduled Time", "Interviewer", "Status", "Rating", "Outcome", "Feedback Comments",
                "Client Feedback Comments",
            ]
            filename = f"interview_status_summary_{today_str}.{payload.attachment_format}"
            if payload.attachment_format == "xlsx":
                msg.add_attachment(
                    build_summary_xlsx(export_headers, export_rows, status_colors),
                    maintype="application",
                    subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    filename=filename,
                )
            else:
                msg.add_attachment(
                    build_summary_csv(export_headers, export_rows),
                    maintype="text",
                    subtype="csv",
                    filename=filename,
                )

        try:
            password = decrypt_password(config_model.smtp_password)
            if config_model.use_ssl:
                server = smtplib.SMTP_SSL(config_model.smtp_server, config_model.smtp_port)
            else:
                server = smtplib.SMTP(config_model.smtp_server, config_model.smtp_port)
                if config_model.use_tls:
                    server.starttls()
            server.login(config_model.smtp_username, password)
            server.send_message(msg, to_addrs=to_addrs)
            server.quit()
        except Exception as e:
            logger.error(f"Failed to send interview summary email: {e}")
            raise HTTPException(status_code=500, detail=f"Failed to send email: {str(e)}")

        logger.info(f"Sent interview summary of {len(docs)} sessions to: {to_addrs}")
        return {
            "status": "success",
            "message": f"Interview summary for {len(docs)} session(s) sent to {', '.join(to_addrs)}",
            "recipients": to_addrs,
            "sent_count": len(docs),
        }

    async def get_feedback_questions(self, interview_type: Optional[str] = None) -> Dict[str, Any]:
        """Load and return feedback questions structured by interview type from JSON file."""
        import json
        from pathlib import Path
        json_path = Path(__file__).resolve().parent.parent / "data" / "feedback_questions.json"
        
        questions_dict = {}
        if json_path.exists():
            with open(json_path, "r", encoding="utf-8") as f:
                questions_dict = json.load(f)

        if not interview_type:
            return questions_dict

        norm_type = interview_type.upper().strip()
        
        if norm_type in questions_dict:
            questions = questions_dict[norm_type]
        elif "TECH" in norm_type or "CODING" in norm_type:
            questions = questions_dict.get("TECHNICAL", questions_dict.get("DEFAULT", []))
        elif "HR" in norm_type or "SCREENING" in norm_type:
            questions = questions_dict.get("HR", questions_dict.get("DEFAULT", []))
        elif "MANAGERIAL" in norm_type or "CULTURE" in norm_type:
            questions = questions_dict.get("MANAGERIAL", questions_dict.get("DEFAULT", []))
        elif "FINAL" in norm_type or "CLIENT" in norm_type:
            questions = questions_dict.get("FINAL_ROUND", questions_dict.get("DEFAULT", []))
        else:
            questions = questions_dict.get("DEFAULT", [])

        return {
            "interview_type": norm_type,
            "questions": questions
        }

    async def check_candidate_active_status(self, candidate_id: str, candidate_name: Optional[str] = None) -> Dict[str, Any]:
        """Check candidate active incomplete interview status directly from database."""
        active_interviews = await self.interview_repo.get_active_incomplete_by_candidate(
            candidate_id=candidate_id, candidate_name=candidate_name
        )
        if active_interviews:
            active = active_interviews[0]
            st = active.get("status", "SCHEDULED")
            cand_name = active.get("candidate_name") or candidate_name or "Candidate"
            return {
                "has_active_interview": True,
                "interview_id": active.get("id"),
                "candidate_id": active.get("candidate_id"),
                "candidate_name": cand_name,
                "status": st,
                "job_title": active.get("job_title"),
                "scheduled_date": active.get("scheduled_date"),
                "scheduled_time": active.get("scheduled_time"),
                "round_number": active.get("round_number", 1),
                "message": f"Cannot assign new interview. Candidate '{cand_name}' already has an assigned interview session (Status: {st}) that is not yet completed. Please complete feedback for the previous round first."
            }
        return {
            "has_active_interview": False,
            "message": "Candidate has no active pending interviews."
        }

    async def cancel_interview(
        self,
        interview_id: str,
        payload: InterviewCancelRequest,
        updated_by: Optional[str] = None,
    ) -> InterviewResponse:
        """Cancel an interview session with optional cancellation reason and email dispatch."""
        existing = await self.interview_repo.get_by_id(interview_id)
        if not existing:
            raise NotFoundError("Interview not found.")

        reason_str = payload.reason.strip() if payload.reason and payload.reason.strip() else None

        update_fields: Dict[str, Any] = {
            "status": InterviewStatus.CANCELLED,
            "updated_by": updated_by,
            "updated_at": utc_now().isoformat(),
        }
        if reason_str:
            existing_notes = existing.get("notes") or ""
            cancellation_note = f"[CANCELLED REASON]: {reason_str}"
            update_fields["notes"] = f"{existing_notes}\n{cancellation_note}".strip() if existing_notes else cancellation_note

        updated_doc = await self.interview_repo.update(interview_id, update_fields)

        if payload.send_email:
            try:
                candidate_email = existing.get("candidate_email")
                interviewer_email = existing.get("interviewer_email")
                email_payload = SendInterviewEmailRequest(
                    send_to_candidate=True if candidate_email else False,
                    candidate_email=candidate_email,
                    send_to_interviewer=True if interviewer_email else False,
                    interviewer_email=interviewer_email,
                    custom_notes=f"Interview Session Cancelled. Reason: {reason_str or 'No reason specified'}",
                )
                await self.send_interview_email(interview_id, email_payload)
            except Exception as e:
                logger.warning(f"Interview cancelled successfully but failed to dispatch cancellation email: {e}")

        logger.info(f"Cancelled interview ID '{interview_id}'")
        return InterviewResponse.model_validate(updated_doc)



